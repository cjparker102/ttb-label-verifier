"""Reads a batch's application data from a CSV or Excel file, in memory."""

import csv
import io
import re
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook

from app.rules.application import Application, BeverageType

MAX_ROWS = 300  # applications per batch
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024  # an .xlsx file is a ZIP; refuse ones that balloon

# Accepted ways of writing each column heading, compared after tidying.
COLUMNS = {
    "file_names": [
        "file names",
        "file name",
        "filenames",
        "filename",
        "files",
        "file",
        "images",
        "image",
        "pictures",
    ],
    "beverage_type": ["beverage type", "type of product", "product type", "type"],
    "brand_name": ["brand name", "brand"],
    "class_type": ["class type", "class or type", "class"],
    "alcohol_content": ["alcohol content", "alcohol", "abv"],
    "net_contents": ["net contents", "net content", "volume"],
    "bottler_name_address": [
        "bottler name address",
        "bottler name and address",
        "bottler or importer name and address",
        "name and address",
        "bottler",
        "importer",
    ],
    "country_of_origin": ["country of origin", "country"],
}
REQUIRED = [
    "file_names",
    "beverage_type",
    "brand_name",
    "class_type",
    "net_contents",
    "bottler_name_address",
]
HEADINGS = {
    "file_names": "file names",
    "beverage_type": "beverage type",
    "brand_name": "brand name",
    "class_type": "class or type",
    "alcohol_content": "alcohol content",
    "net_contents": "net contents",
    "bottler_name_address": "bottler name and address",
    "country_of_origin": "country of origin",
}


class SpreadsheetError(ValueError):
    """The file can't be used as a batch spreadsheet. The message says why, in plain words."""


@dataclass(frozen=True)
class ApplicationRow:
    row_number: int  # as numbered in the spreadsheet, where the headings are row 1
    file_names: tuple[str, ...]
    application: Application | None  # None when the row has problems
    problems: tuple[str, ...]
    brand_name: str = ""  # as typed, so a row with problems can still be recognized


def tidy_heading(text: str) -> str:
    return " ".join(re.sub(r"[_\-/]+", " ", text).casefold().split())


ALIASES = {tidy_heading(alias): column for column, aliases in COLUMNS.items() for alias in aliases}


def split_file_names(cell: str) -> tuple[str, ...]:
    return tuple(name.strip() for name in re.split(r"[;\n|]", cell) if name.strip())


def read_spreadsheet(name: str, data: bytes) -> list[ApplicationRow]:
    extension = name.rsplit(".", 1)[-1].casefold() if "." in name else ""
    if extension == "csv":
        table = csv_rows(data)
    elif extension == "xlsx":
        table = xlsx_rows(data)
    else:
        raise SpreadsheetError(
            f'"{name}" is not a spreadsheet this app can read. Use a CSV or Excel (.xlsx) file.'
        )
    return application_rows(table)


def csv_rows(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252"):  # cp1252: what Excel on Windows saves
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SpreadsheetError("This CSV file could not be read. Save it as UTF-8 and try again.")
    try:
        dialect = csv.Sniffer().sniff(text.split("\n", 1)[0], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def xlsx_rows(data: bytes) -> list[list[str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            unzipped = sum(entry.file_size for entry in package.infolist())
    except zipfile.BadZipFile as error:
        raise SpreadsheetError("This Excel file could not be opened.") from error
    if unzipped > MAX_UNZIPPED_BYTES:
        raise SpreadsheetError("This Excel file is too large once opened.")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise SpreadsheetError("This Excel file could not be opened.") from error
    try:
        sheet = workbook.worksheets[0]
        return [[cell_text(cell) for cell in row] for row in sheet.iter_rows()]
    finally:
        workbook.close()


def cell_text(cell) -> str:
    """Excel stores 13.5% as 0.135 and 45 as 45.0; turn values back into what people typed."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        if "%" in (getattr(cell, "number_format", "") or ""):
            value = round(value * 100, 6)
            return f"{value:g}%"
        return f"{value:g}"
    return str(value)


def application_rows(table: list[list[str]]) -> list[ApplicationRow]:
    if not table:
        raise SpreadsheetError("The spreadsheet is empty.")
    columns: dict[str, int] = {}
    for index, heading in enumerate(table[0]):
        column = ALIASES.get(tidy_heading(heading))
        if column and column not in columns:
            columns[column] = index
    missing = [HEADINGS[c] for c in REQUIRED if c not in columns]
    if missing:
        raise SpreadsheetError(
            "The spreadsheet needs a column for: "
            + ", ".join(missing)
            + ". Check the headings in row 1."
        )

    rows = []
    for number, cells in enumerate(table[1:], start=2):
        values = {c: (cells[i].strip() if i < len(cells) else "") for c, i in columns.items()}
        if not any(values.values()):
            continue  # blank row
        rows.append(build_row(number, values))
    if not rows:
        raise SpreadsheetError("The spreadsheet has headings but no applications.")
    if len(rows) > MAX_ROWS:
        raise SpreadsheetError(
            f"The spreadsheet has {len(rows)} applications. One batch can have up to {MAX_ROWS}."
        )
    return rows


def build_row(number: int, values: dict[str, str]) -> ApplicationRow:
    problems = []
    file_names = split_file_names(values["file_names"])
    if not file_names:
        problems.append("No picture file names are listed.")
    beverage_type = None
    try:
        beverage_type = BeverageType.from_text(values["beverage_type"])
    except ValueError:
        shown = values["beverage_type"] or "(blank)"
        problems.append(
            f'Unknown beverage type "{shown}". Use distilled spirits, wine, or malt beverage.'
        )
    for column in REQUIRED[2:]:
        if not values[column]:
            problems.append(f"The {HEADINGS[column]} is blank.")
    application = None
    if not problems:
        application = Application(
            beverage_type=beverage_type,
            brand_name=values["brand_name"],
            class_type=values["class_type"],
            alcohol_content=values.get("alcohol_content", ""),
            net_contents=values["net_contents"],
            bottler_name_address=values["bottler_name_address"],
            country_of_origin=values.get("country_of_origin", ""),
        )
    return ApplicationRow(number, file_names, application, tuple(problems), values["brand_name"])
