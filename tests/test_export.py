import asyncio
from io import BytesIO

from openpyxl import load_workbook

from app.batch import new_job, run_job
from app.export import results_workbook
from app.intake.matching import MatchedRow
from app.intake.spreadsheet import ApplicationRow
from app.rules.application import Application, BeverageType
from tests.disk_watch import watch_disk_writes
from tests.stand_ins import ReplayReader
from tools.make_samples import LABELS_DIR

FRONT = (LABELS_DIR / "old-tom-bourbon_front.png").read_bytes()


def application(brand="OLD TOM DISTILLERY", alcohol="45% Alc./Vol. (90 Proof)"):
    return Application(
        BeverageType.DISTILLED_SPIRITS,
        brand,
        "Kentucky Straight Bourbon Whiskey",
        alcohol,
        "750 mL",
        "Old Tom Distillery, Bardstown, Kentucky",
    )


def matched(number, app, brand, pictures, problems=()):
    names = tuple(name for name, _ in pictures) or ("gone.png",)
    return MatchedRow(ApplicationRow(number, names, app, problems, brand), pictures, ())


def finished_job():
    rows = [
        matched(2, application(), "OLD TOM DISTILLERY", (("front.png", FRONT),)),
        matched(3, application(alcohol="40%"), "OLD TOM DISTILLERY", (("front.png", FRONT),)),
        # A brand that looks like a formula, as a mischievous spreadsheet might contain.
        matched(4, None, '=HYPERLINK("x")', (), ("The brand name is blank.",)),
    ]
    job = new_job(rows, [], [])
    asyncio.run(run_job(job, ReplayReader(), 1568, 2, asyncio.Semaphore(1)))
    return job


def open_workbook(data: bytes):
    return load_workbook(BytesIO(data))


def test_summary_has_one_row_per_label_in_spreadsheet_order():
    sheet = open_workbook(results_workbook(finished_job()))["Summary"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (
        "Row",
        "Brand (application)",
        "Pictures",
        "Overall result",
        "What to look at",
        "Read by",
    )
    assert [r[0] for r in rows[1:]] == [2, 3, 4]
    assert rows[1][3] == "All good"
    assert rows[1][5] == "AI reader"
    assert len(rows[1]) == 6  # no time per label: in a batch it includes waiting
    assert rows[2][3] == "Problems found (1)"
    assert "Alcohol content: The label says 45% but the application says 40%." in rows[2][4]
    assert rows[3][3] == "Not checked"
    assert rows[3][4] == "The brand name is blank."


def test_details_have_one_row_per_field():
    sheet = open_workbook(results_workbook(finished_job()))["Details"]
    rows = list(sheet.iter_rows(values_only=True))[1:]
    assert len(rows) == 14  # 7 fields for each of the 2 labels that were checked
    alcohol = next(r for r in rows if r[0] == 3 and r[2] == "Alcohol content")
    assert alcohol[3:6] == ("40%", "45% Alc./Vol. (90 Proof)", "Problem")
    warning = next(r for r in rows if r[0] == 2 and r[2] == "Government warning")
    assert warning[3] == "The required warning text"
    assert "27 CFR 16.21" in warning[7]


def test_text_that_looks_like_a_formula_stays_text():
    workbook = open_workbook(results_workbook(finished_job()))
    cell = workbook["Summary"]["B4"]
    assert cell.value == '=HYPERLINK("x")'
    assert cell.data_type == "s"


def test_building_the_workbook_writes_nothing_to_disk():
    job = finished_job()
    results_workbook(job)  # warm-up: loads any lazy imports
    with watch_disk_writes() as writes:
        results_workbook(job)
    assert list(writes) == []


def test_label_text_that_looks_like_a_formula_stays_text():
    rows = [
        MatchedRow(
            ApplicationRow(2, ("front.png",), application(brand="=1+1"), (), "=1+1"),
            (("front.png", FRONT),),
            (),
        )
    ]
    job = new_job(rows, [], [])
    asyncio.run(run_job(job, ReplayReader(), 1568, 2, asyncio.Semaphore(1)))
    details = open_workbook(results_workbook(job))["Details"]
    brand_cells = [row[1] for row in details.iter_rows(min_row=2)]
    assert all(cell.value == "=1+1" and cell.data_type == "s" for cell in brand_cells)
