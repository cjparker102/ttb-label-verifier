"""Batch upload, end to end, with stand-ins for Claude."""

import asyncio
import io
import json
import re
import threading
import time
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.intake.spreadsheet import read_spreadsheet
from app.limits import Limits
from app.main import app, get_reader
from tests.disk_watch import watch_disk_writes
from tests.stand_ins import PerfectReader, ReplayReader, SlowReader, expected_word
from tools.make_samples import LABELS_DIR, SAMPLES_DIR

EXPECTED = json.loads((SAMPLES_DIR / "expected.json").read_text(encoding="utf-8"))
SPREADSHEET = (SAMPLES_DIR / "applications.csv").read_bytes()


def planted_counts() -> tuple[int, int, int]:
    """How many sample labels should come out with problems, needing review, or neither."""
    worst = []
    for entry in EXPECTED:
        results = set(entry["expected"].values())
        worst.append("fail" if "fail" in results else "review" if "review" in results else "pass")
    return worst.count("fail"), worst.count("review"), worst.count("pass")


def samples_zip(extra=()) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for entry in EXPECTED:
            for name in entry["images"]:
                archive.write(LABELS_DIR / name, f"labels/{name}")
        for name, data in extra:
            archive.writestr(name, data)
    return out.getvalue()


@pytest.fixture
def client():
    saved = app.state.batch_concurrency
    with TestClient(app) as test_client:  # keeps the event loop running for background batches
        yield test_client
    app.dependency_overrides.clear()
    app.state.batch_concurrency = saved


def use_reader(reader):
    app.dependency_overrides[get_reader] = lambda: reader


def start(client, spreadsheet=SPREADSHEET, pictures=None, name="labels.zip"):
    files = [
        ("spreadsheet", ("applications.csv", spreadsheet, "text/csv")),
        (
            "pictures",
            (name, pictures if pictures is not None else samples_zip(), "application/zip"),
        ),
    ]
    return client.post("/batches", files=files, follow_redirects=False)


# The sample batch reads one label at a time here, and each label's warning picture is read a
# second time by Tesseract: about 14 s on a laptop, and over 30 s on a slow CI runner.
def wait_until_finished(client, url, timeout=120):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        progress = client.get(f"{url}/progress").json()
        if progress["state"] == "finished":
            return progress
        time.sleep(0.05)
    raise AssertionError("the batch did not finish")


def run_samples_batch(client, extra=()):
    use_reader(PerfectReader([entry["truth"] for entry in EXPECTED]))
    app.state.batch_concurrency = 1  # one at a time, so readings come back in row order
    response = start(client, pictures=samples_zip(extra))
    assert response.status_code == 303
    url = response.headers["location"]
    wait_until_finished(client, url)
    return url


def test_every_sample_in_a_batch_gets_its_expected_results(client):
    url = run_samples_batch(client)
    job = app.state.batches.get(url.rsplit("/", 1)[-1])
    items = sorted(job.items, key=lambda item: item.row_number)
    assert len(items) == len(EXPECTED)
    for item, entry in zip(items, EXPECTED, strict=True):
        got = {field.field: expected_word(field) for field in item.outcome.result.fields}
        expected = dict(entry["expected"])
        if entry.get("photo") and got["government_warning"] == "pass_with_note":
            # On some photo-like labels code can't find or measure the heading, so the AI
            # reader's "bold" is not confirmed: still Matches, with a note.
            expected["government_warning"] = {"pass": "pass_with_note"}.get(
                expected["government_warning"]
            )
        assert got == expected, entry["id"]


def test_results_page_shows_problems_first_and_the_counts(client):
    url = run_samples_batch(client, extra=[("readme.txt", b"notes"), (".DS_Store", b"x")])
    page = client.get(url).text
    assert f"{len(EXPECTED)} of {len(EXPECTED)} labels checked." in page
    fails, reviews, passes = planted_counts()
    assert f"Problems found in {fails} labels" in page
    assert f"{reviews} labels need review" in page
    assert f"{passes} labels with no issues found in the checked fields" in page
    assert "Download results (Excel)" in page
    assert page.index("RIDGEWICK") < page.index("TIDEWICK")  # a problem comes before a pass
    # The batch's time, not each label's: that would include waiting for other labels.
    assert "Checked in " not in page
    # A busy test machine can take over a minute: "1 minute 4 seconds".
    took = r"(?:\d+ minutes? )?\d+ seconds?|\d+ minutes?"
    assert re.search(
        rf"Checked {len(EXPECTED)} labels in (?:{took}), about [\d.]+ seconds a label\.", page
    )
    assert "labels.zip: readme.txt" in page
    assert "labels.zip: .DS_Store" in page


def test_excel_download_has_every_label(client):
    url = run_samples_batch(client)
    response = client.get(f"{url}/results.xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    summary = load_workbook(io.BytesIO(response.content))["Summary"]
    assert summary.max_row == len(EXPECTED) + 1  # the headings, then one row per label


def test_previews_are_served_from_memory(client):
    url = run_samples_batch(client)
    preview = client.get(f"{url}/previews/2/1.jpg")
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "image/jpeg"
    assert client.get(f"{url}/previews/99/1.jpg").status_code == 404
    assert client.get(f"{url}/previews/2/9.jpg").status_code == 404
    assert client.get(f"{url}/previews/2/0.jpg").status_code == 404


def test_each_label_has_a_full_width_row_with_its_reasons_fields_and_pictures(client):
    url = run_samples_batch(client)
    page = client.get(url).text
    assert page.count('<tr class="detail-row') == len(EXPECTED)
    assert '<td colspan="3">' in page
    assert "Show every field and picture for row 2" in page
    assert f'src="{url}/previews/2/1.jpg"' in page
    assert 'alt="Picture 1 of 2 for row 2"' in page  # the sample has a front and a back
    assert f'src="{url}/previews/2/2.jpg"' in page


def test_while_running_the_page_says_so_and_excel_waits(client):
    hold = threading.Event()

    class HeldReader(ReplayReader):
        async def read(self, images):
            while not hold.is_set():
                await asyncio.sleep(0.01)
            return await super().read(images)

    use_reader(HeldReader())
    url = start(client).headers["location"]
    try:
        page = client.get(url).text
        assert "This page updates by itself." in page
        assert '<meta http-equiv="refresh" content="5">' in page  # for browsers without scripts
        assert client.get(f"{url}/results.xlsx").status_code == 409
        assert client.get(f"{url}/progress").json()["state"] in ("queued", "running")
    finally:
        hold.set()
    wait_until_finished(client, url)


def test_rows_whose_pictures_are_missing_are_reported(client):
    use_reader(ReplayReader())
    # The spreadsheet names a picture that isn't in the ZIP, and so leaves one unused.
    sheet = SPREADSHEET.replace(b"stones-throw-gin.png", b"missing.png")
    url = start(client, spreadsheet=sheet).headers["location"]
    wait_until_finished(client, url)
    page = client.get(url).text
    assert "These pictures were not in the upload: missing.png." in page
    assert "1 label could not be checked" in page
    assert "<strong>stones-throw-gin.png</strong>: No row in the spreadsheet lists" in page


def test_spreadsheet_template_downloads_and_reads_back(client):
    response = client.get("/template.csv")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    rows = read_spreadsheet("template.csv", response.content)
    assert [row.problems for row in rows] == [(), ()]


@pytest.mark.parametrize(
    ("spreadsheet", "pictures", "message"),
    [
        (b"brand\nOLD TOM\n", None, "needs a column for"),
        (SPREADSHEET, b"not a zip", "could not be opened"),
    ],
)
def test_unusable_uploads_say_why(client, spreadsheet, pictures, message):
    use_reader(ReplayReader())
    response = start(client, spreadsheet=spreadsheet, pictures=pictures)
    assert response.status_code == 422
    assert message in response.text


NO_SPREADSHEET = "Add the spreadsheet of applications, or put it in the ZIP file with the pictures."


def test_missing_spreadsheet_or_pictures(client):
    use_reader(ReplayReader())
    response = client.post(
        "/batches", files=[("pictures", ("a.zip", samples_zip(), "application/zip"))]
    )
    assert response.status_code == 422
    assert NO_SPREADSHEET in response.text  # none in the box, and none in the ZIP file
    response = client.post("/batches", files=[("spreadsheet", ("a.csv", SPREADSHEET, "text/csv"))])
    assert "Add the label pictures, or a ZIP file of them." in response.text
    assert NO_SPREADSHEET not in response.text
    response = client.post("/batches", files=[("other", ("x.txt", b"x", "text/plain"))])
    assert NO_SPREADSHEET in response.text
    assert "Add the label pictures, or a ZIP file of them." in response.text


def test_the_spreadsheet_box_may_be_left_empty(client):
    page = client.get("/").text
    box = re.search(r'<input type="file" id="spreadsheet"[^>]*>', page).group(0)
    assert "required" not in box  # the browser lets the form go with only a ZIP file
    assert "Or leave this box empty and put the spreadsheet in the ZIP file" in page


def notes_on(client, response) -> list[str]:
    assert response.status_code == 303
    url = response.headers["location"]
    wait_until_finished(client, url)
    return app.state.batches.get(url.rsplit("/", 1)[-1]).notes


def test_a_spreadsheet_inside_the_zip_is_used_when_the_box_is_empty(client):
    use_reader(ReplayReader())
    archive = samples_zip(extra=[("batch/applications.csv", SPREADSHEET)])
    response = client.post(
        "/batches",
        files=[("pictures", ("labels.zip", archive, "application/zip"))],
        follow_redirects=False,
    )
    job = app.state.batches.get(response.headers["location"].rsplit("/", 1)[-1])
    assert notes_on(client, response) == [
        'Spreadsheet used: "batch/applications.csv" in labels.zip.'
    ]
    assert job.total == len(EXPECTED)
    assert all(item.outcome is not None for item in job.items)
    assert job.skipped_files == []  # the spreadsheet isn't listed as a skipped picture
    page = " ".join(client.get(response.headers["location"]).text.split())
    assert "Spreadsheet used: &#34;batch/applications.csv&#34; in labels.zip." in page


def test_with_several_spreadsheets_in_the_zip_the_first_by_name_is_used(client):
    use_reader(ReplayReader())
    one_row = SPREADSHEET.split(b"\n", 2)
    archive = samples_zip(
        extra=[("b/more.csv", one_row[0] + b"\n" + one_row[1]), ("a/all.csv", SPREADSHEET)]
    )
    response = start_with_pictures_only(client, archive)
    assert notes_on(client, response) == [
        'Spreadsheet used: "a/all.csv" in labels.zip, the first by name of the 2 found. '
        'Not used: "b/more.csv" in labels.zip.'
    ]
    job = app.state.batches.get(response.headers["location"].rsplit("/", 1)[-1])
    assert job.total == len(EXPECTED)  # every row of a/all.csv, not the one row of b/more.csv


def test_the_spreadsheet_in_the_box_wins_over_one_in_the_zip(client):
    use_reader(ReplayReader())
    one_row = b"\n".join(SPREADSHEET.split(b"\n", 2)[:2])
    archive = samples_zip(extra=[("applications.csv", SPREADSHEET)])
    response = start(client, spreadsheet=one_row, pictures=archive)
    assert notes_on(client, response) == [
        'Spreadsheet used: "applications.csv", the one you added. '
        'Not used: "applications.csv" in labels.zip.'
    ]
    job = app.state.batches.get(response.headers["location"].rsplit("/", 1)[-1])
    assert job.total == 1


def test_a_broken_spreadsheet_in_the_zip_is_named(client):
    use_reader(ReplayReader())
    archive = samples_zip(extra=[("applications.csv", b"just,some\nwords,here\n")])
    response = start_with_pictures_only(client, archive)
    assert response.status_code == 422
    assert "The spreadsheet needs a column for" in response.text
    assert "Spreadsheet used: &#34;applications.csv&#34; in labels.zip." in response.text


def start_with_pictures_only(client, archive):
    return client.post(
        "/batches",
        files=[("pictures", ("labels.zip", archive, "application/zip"))],
        follow_redirects=False,
    )


def test_batch_limits_are_enforced(client):
    use_reader(ReplayReader())
    app.state.limits = Limits(1000, 600, 0, 3600, 100_000)  # no batches allowed
    response = start(client)
    assert response.status_code == 429
    assert "batches" in response.text


def test_expired_or_unknown_batches(client):
    assert client.get("/batches/no-such-batch").status_code == 404
    assert "no longer available" in client.get("/batches/no-such-batch").text
    assert client.get("/batches/no-such-batch/progress").json()["state"] == "gone"


def test_a_whole_batch_writes_nothing_to_disk(client):
    url = run_samples_batch(client)  # warm-up: loads any lazy imports
    client.get(f"{url}/results.xlsx")
    with watch_disk_writes() as writes:
        url = run_samples_batch(client)
        client.get(url)
        client.get(f"{url}/results.xlsx")
    assert list(writes) == []


def test_a_batch_sent_as_a_7z_file(client):
    from tests.intake.test_sevenzip import make_7z

    use_reader(ReplayReader())
    pictures = [(name, (LABELS_DIR / name).read_bytes()) for e in EXPECTED for name in e["images"]]
    response = start(client, pictures=make_7z(*pictures), name="labels.7z")
    assert response.status_code == 303
    progress = wait_until_finished(client, response.headers["location"])
    assert progress["done"] == progress["total"] == len(EXPECTED)


class HeldReader(ReplayReader):
    """Holds every reading until it is let go, so the pictures stay waiting in memory."""

    def __init__(self):
        super().__init__()
        self.let_go = False

    async def read(self, images):
        while not self.let_go:
            await asyncio.sleep(0.01)
        return await super().read(images)


SAMPLE_PICTURE_BYTES = sum((LABELS_DIR / n).stat().st_size for e in EXPECTED for n in e["images"])


def test_batches_are_turned_away_while_waiting_pictures_would_pass_the_cap(client, monkeypatch):
    reader = HeldReader()
    use_reader(reader)
    # Room for the pictures of two sample batches, not three.
    monkeypatch.setattr(app.state, "max_waiting_bytes", int(2.5 * SAMPLE_PICTURE_BYTES))
    running, waiting = start(client), start(client)
    assert (running.status_code, waiting.status_code) == (303, 303)
    turned_away = start(client)
    assert turned_away.status_code == 503
    page = " ".join(turned_away.text.split())
    assert "The app is busy with other batches. Please try again in 1 minute." in page
    assert app.state.limits.daily.used == 2 * len(EXPECTED)  # the third wasn't counted

    reader.let_go = True
    for response in (running, waiting):
        wait_until_finished(client, response.headers["location"])
    after = start(client)  # checked labels let go of their pictures, so there is room again
    assert after.status_code == 303
    wait_until_finished(client, after.headers["location"])


def test_a_batch_bigger_than_the_cap_is_turned_away_at_once(client, monkeypatch):
    use_reader(ReplayReader())
    monkeypatch.setattr(app.state, "max_waiting_bytes", SAMPLE_PICTURE_BYTES // 2)
    response = start(client)
    assert response.status_code == 413
    assert "Please split it into smaller batches." in response.text
    assert app.state.limits.daily.used == 0


def test_batch_labels_are_never_asked_for_twice(client, monkeypatch):
    monkeypatch.setattr(app.state, "hedge_after", 0.01)
    reader = SlowReader(0.05)  # every answer is slower than the wait before asking again
    use_reader(reader)
    response = start(client)
    progress = wait_until_finished(client, response.headers["location"])
    assert reader.requests == progress["total"] == len(EXPECTED)


def test_progress_of_an_expired_batch_carries_the_message_the_page_shows(client):
    response = client.get("/batches/no-such-batch/progress")
    assert response.status_code == 404
    answer = response.json()
    assert answer["state"] == "gone"
    assert (
        "no longer available" in answer["message"] and "upload the batch again" in answer["message"]
    )


def test_the_batch_page_has_a_place_to_say_it_is_reconnecting(client):
    url = run_samples_batch(client)
    page = client.get(url).text
    assert 'id="batch-connection" role="status" hidden' in page
