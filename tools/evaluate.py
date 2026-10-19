"""Measures how well the app reads and judges the sample labels, and how fast.

Two ways to run it, from the project root:

  In this process, with the real Claude reader (the key is read from .env):
    uv run --env-file .env python -m tools.evaluate

  Against a deployed app, end to end through its web pages:
    uv run python -m tools.evaluate --url https://... --deployed-as "revision ..., image ..."
  The app lets one visitor make only so many checks in ten minutes. When a run reaches
  that limit, it waits a minute and tries again; the wait isn't counted in the times.

Each run reads every sample label once (31 labels as of Oct 19, 2026), which costs
about two cents with Claude Haiku 5.5. Results are saved to docs/measurements/ as
JSON and Markdown, exactly as measured.

How results are judged, field by field, against samples/expected.json:
  exact                  the expected result
  false pass             a planted problem shown as Matches: the error that matters most
  false fail             a Problem where none was planted (or only "Check this" was)
  review instead of fail a planted problem flagged as "Check this" rather than a Problem
  extra review           "Check this" on a field that is fine
On clean labels only "exact" is right. On photo-like labels, "Check this" is also an
acceptable answer where the photo is hard to read.
"""

import argparse
import asyncio
import io
import json
import math
import re
import statistics
import subprocess
import time
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from app.rules.application import Application, BeverageType
from tools.make_samples import LABELS_DIR, SAMPLES_DIR, sample_for

OUT_DIR = Path(__file__).resolve().parent.parent / "docs" / "measurements"
PASS_WORDS = {"pass", "pass_with_note"}
# Claude Haiku 5.5, USD per million tokens, for prompts up to 100K tokens.
PRICE_IN, PRICE_OUT = 0.10, 0.50
WORDS = {"Matches": "pass", "Check this": "review", "Problem": "fail", "Not needed": "not_needed"}
TITLES = {
    "Brand name": "brand_name",
    "Class / type": "class_type",
    "Alcohol content": "alcohol_content",
    "Net contents": "net_contents",
    "Bottler name and address": "bottler_name_address",
    "Country of origin": "country_of_origin",
    "Government warning": "government_warning",
}


def result_word(field) -> str:
    """A field result in the words samples/expected.json uses."""
    if field.status == "pass" and field.notes:
        return "pass_with_note"
    return str(field.status)


def judge(expected: str, got: str) -> str:
    """How one field's result compares with what was planted."""
    if got == expected:
        return "exact"
    if expected in ("fail", "review") and got in PASS_WORDS | {"not_needed"}:
        return "false_pass"
    if got == "fail":
        return "false_fail"
    if got == "review":
        return "review_instead_of_fail" if expected == "fail" else "extra_review"
    if expected in PASS_WORDS and got in PASS_WORDS:
        return "note_difference"
    return "other"


ACCEPTABLE_ON_PHOTOS = {"exact", "note_difference", "review_instead_of_fail", "extra_review"}


def percentile(values: list[float], share: float) -> float:
    """Nearest-rank percentile, so the number is one that was really measured."""
    ordered = sorted(values)
    return ordered[max(0, math.ceil(share * len(ordered)) - 1)]


def load_entries(only: list[str] | None) -> list[dict]:
    entries = json.loads((SAMPLES_DIR / "expected.json").read_text(encoding="utf-8"))
    return [e for e in entries if not only or e["id"] in only]


def application_for(entry: dict) -> Application:
    sample = sample_for(entry["id"])
    return Application(BeverageType.from_text(sample.beverage_type), **sample.application)


def score(entry: dict, got: dict[str, str]) -> dict:
    judgments = {name: judge(entry["expected"][name], got[name]) for name in entry["expected"]}
    return {"got": got, "judgments": judgments}


# In this process


async def evaluate_here(entries: list[dict], reader_name: str) -> dict:
    import anthropic

    from app.checking import NoReader, check_one
    from app.config import load_settings
    from app.reading.claude import ClaudeReader
    from tools.time_reading import compare

    settings = load_settings()
    reader = None
    if reader_name == "anthropic":
        client = anthropic.AsyncAnthropic(max_retries=1, timeout=60)
        reader = ClaudeReader(
            client, thinking=settings.thinking, effort=settings.effort, model=settings.model
        )

    records = []
    for entry in entries:
        images = [(name, (LABELS_DIR / name).read_bytes()) for name in entry["images"]]
        record = {"id": entry["id"], "photo": entry.get("photo")}
        try:
            outcome = await check_one(
                application_for(entry), images, reader, settings.max_image_edge,
                ocr_fallback=reader is None, outline=True,  # as single checks in the app
            )  # fmt: skip
        except NoReader as error:
            record["error"] = f"{type(error).__name__}: {error.__cause__ or error}"
            print(f"{entry['id']:50} ERROR {record['error']}")
            records.append(record)
            continue
        got = {field.field: result_word(field) for field in outcome.result.fields}
        reading_problems = (
            compare(outcome.reading, entry["truth"]) if outcome.read_by == "ai" else {}
        )
        record.update(
            score(entry, got),
            seconds=round(outcome.seconds, 2),
            read_by=outcome.read_by,
            input_tokens=outcome.input_tokens,
            output_tokens=outcome.output_tokens,
            reading_problems=reading_problems,
            reading=outcome.reading.model_dump(mode="json"),
            # Fields to look at that are printed on the label, and those outlined.
            flagged_on_label=[
                f.field
                for f in outcome.result.fields
                if f.status in ("fail", "review") and f.label_says
            ],
            outlined=[highlight.field for highlight in outcome.highlights],
        )
        print(f"{entry['id']:50} {outcome.seconds:5.2f} s  {describe(record)}")
        records.append(record)
    return {
        "mode": "in process",
        "reader": reader_name,
        "settings": {
            "model": settings.model,
            "thinking": "adaptive" if settings.thinking else "off",
            "effort": settings.effort,
            "max_image_edge": settings.max_image_edge,
        },
        "records": records,
    }


# Against a deployed app


def samples_zip(entries: list[dict]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for entry in entries:
            for name in entry["images"]:
                archive.write(LABELS_DIR / name, name)
    return out.getvalue()


def spreadsheet(entries: list[dict]) -> bytes:
    import csv

    out = io.StringIO()
    writer = csv.writer(out)
    columns = ["file_names", "beverage_type", "brand_name", "class_type", "alcohol_content",
               "net_contents", "bottler_name_address", "country_of_origin"]  # fmt: skip
    writer.writerow(columns)
    for entry in entries:
        application = sample_for(entry["id"]).application
        writer.writerow(
            [
                "; ".join(entry["images"]),
                entry["beverage_type"],
                *(application[c] for c in columns[2:]),
            ]
        )
    return out.getvalue().encode("utf-8")


def results_from_excel(data: bytes) -> dict[int, dict]:
    """Field results and reader by spreadsheet row, from the app's Excel download."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data))
    rows: dict[int, dict] = {}
    for row in workbook["Summary"].iter_rows(min_row=2, values_only=True):
        rows[row[0]] = {"overall": row[3], "read_by": row[5], "got": {}}
    for row in workbook["Details"].iter_rows(min_row=2, values_only=True):
        number, title, word = row[0], row[2], row[5]
        got = "pass_with_note" if word.startswith("Matches,") else WORDS[word]
        rows[number]["got"][TITLES[title]] = got
    return rows


FLAGGED_ROW = re.compile(r'<tr class="row-(?:fail|review)">(.*?)</tr>', re.DOTALL)
ROW_TITLE = re.compile(r'<th scope="row">(.*?)</th>')
CHECKED_IN = re.compile(r"Checked in ([\d.]+) seconds")


def from_page(page: str) -> dict:
    """What a single-check result page shows: who read the label, how long the check
    took inside the app, and which fields to look at are printed on the label and
    outlined on the picture."""
    flagged, outlined = [], []
    for row in FLAGGED_ROW.findall(page):
        field = TITLES[ROW_TITLE.search(row).group(1)]
        if 'class="label-text' in row:
            flagged.append(field)
        if "on the picture</span>" in row:
            outlined.append(field)
    checked_in = CHECKED_IN.search(page)
    return {
        "read_by": "ocr" if "read with basic text reading" in page else "ai",
        "server_seconds": float(checked_in.group(1)) if checked_in else None,
        "flagged_on_label": flagged,
        "outlined": outlined,
    }


LIMIT_WAIT = 60  # seconds


def post_check(client, url: str, form: dict, files: list):
    """Post one single-label check, waiting whenever the app's limit on checks per
    visitor is reached. Returns the response and how long it took, without the waits."""
    for _ in range(15):
        start = time.perf_counter()
        response = client.post(f"{url}/check", data=form, files=files)
        seconds = time.perf_counter() - start
        if response.status_code != 429:
            return response, seconds
        print(f"The app's limit on checks was reached; waiting {LIMIT_WAIT} s")
        time.sleep(LIMIT_WAIT)
    raise SystemExit("The app kept refusing checks: HTTP 429")


def check_form(entry: dict) -> dict[str, str]:
    """A sample's application, as the single-check form sends it."""
    application = application_for(entry)
    return {
        "beverage_type": application.beverage_type.value,
        "brand_name": application.brand_name,
        "class_type": application.class_type,
        "alcohol_content": application.alcohol_content,
        "net_contents": application.net_contents,
        "bottler_name_address": application.bottler_name_address,
        "country_of_origin": application.country_of_origin,
    }


def evaluate_deployed(entries: list[dict], url: str) -> dict:
    import httpx

    url = url.rstrip("/")
    records = []
    with httpx.Client(timeout=120) as client:
        for entry in entries:
            form = check_form(entry)
            files = [("images", (n, (LABELS_DIR / n).read_bytes())) for n in entry["images"]]
            response, seconds = post_check(client, url, form, files)
            record = {
                "id": entry["id"],
                "photo": entry.get("photo"),
                "seconds": round(seconds, 2),
                "status_code": response.status_code,
                **from_page(response.text),
            }
            print(f"{entry['id']:50} {seconds:5.2f} s  HTTP {response.status_code}")
            records.append(record)

        start = time.perf_counter()
        response = client.post(
            f"{url}/batches",
            files=[
                ("spreadsheet", ("applications.csv", spreadsheet(entries), "text/csv")),
                ("pictures", ("labels.zip", samples_zip(entries), "application/zip")),
            ],
            follow_redirects=False,
        )
        if response.status_code != 303:
            raise SystemExit(f"The batch upload failed: HTTP {response.status_code}")
        batch = url + response.headers["location"]
        while client.get(f"{batch}/progress").json()["state"] != "finished":
            time.sleep(0.5)
        batch_seconds = time.perf_counter() - start
        excel = client.get(f"{batch}/results.xlsx")
        by_row = results_from_excel(excel.content)

    for number, (entry, record) in enumerate(zip(entries, records, strict=True), start=2):
        record.update(score(entry, by_row[number]["got"]))
        record["batch_read_by"] = by_row[number]["read_by"]
        print(f"{entry['id']:50} {describe(record)}")
    return {
        "mode": "deployed app",
        "url": url,
        "batch_seconds": round(batch_seconds, 1),
        "records": records,
    }


# Reporting


def describe(record: dict) -> str:
    counts = Counter(record["judgments"].values())
    if set(counts) <= {"exact"}:
        return "all fields as expected"
    return ", ".join(
        f"{n} {kind.replace('_', ' ')}" for kind, n in counts.items() if kind != "exact"
    )


def summarize(results: dict) -> dict:
    scored = [r for r in results["records"] if "judgments" in r]
    clean = [r for r in scored if not r["photo"]]
    photos = [r for r in scored if r["photo"]]
    all_judgments = Counter(j for r in scored for j in r["judgments"].values())
    false_passes = [
        f"{r['id']}: {field}"
        for r in scored
        for field, j in r["judgments"].items()
        if j == "false_pass"
    ]
    false_fails = [
        f"{r['id']}: {field}"
        for r in scored
        for field, j in r["judgments"].items()
        if j == "false_fail"
    ]
    seconds = [r["seconds"] for r in results["records"] if "seconds" in r]
    summary = {
        "labels": len(results["records"]),
        "errors": len(results["records"]) - len(scored),
        "fields": sum(all_judgments.values()),
        "judgments": dict(all_judgments),
        "clean_labels_all_exact": sum(
            1 for r in clean if set(r["judgments"].values()) == {"exact"}
        ),
        "clean_labels": len(clean),
        "photo_fields_acceptable": sum(
            1 for r in photos for j in r["judgments"].values() if j in ACCEPTABLE_ON_PHOTOS
        ),
        "photo_fields": sum(len(r["judgments"]) for r in photos),
        "false_passes": false_passes,
        "false_fails": false_fails,
        "read_by_ocr": sum(1 for r in results["records"] if r.get("read_by") == "ocr"),
    }
    if seconds:
        summary["seconds"] = {
            "median": round(statistics.median(seconds), 2),
            "p95": round(percentile(seconds, 0.95), 2),
            "max": round(max(seconds), 2),
            "under_5_seconds": sum(1 for s in seconds if s < 5),
            "timed": len(seconds),
        }
    inside = [r["server_seconds"] for r in results["records"] if r.get("server_seconds")]
    if inside:
        summary["server_seconds"] = {
            "median": round(statistics.median(inside), 2),
            "p95": round(percentile(inside, 0.95), 2),
            "max": round(max(inside), 2),
        }
    reading = [r for r in clean if "reading_problems" in r and r.get("read_by") == "ai"]
    if reading:
        summary["clean_fields_read_exactly"] = sum(7 - len(r["reading_problems"]) for r in reading)
        summary["clean_fields_read"] = 7 * len(reading)
    if any("outlined" in r for r in scored):
        summary["outlined"] = sum(len(r.get("outlined", [])) for r in scored)
        summary["flagged_on_label"] = sum(len(r.get("flagged_on_label", [])) for r in scored)
    tokens_in = sum(r.get("input_tokens", 0) for r in scored)
    tokens_out = sum(r.get("output_tokens", 0) for r in scored)
    if tokens_in:
        summary["cost_usd"] = round((tokens_in * PRICE_IN + tokens_out * PRICE_OUT) / 1e6, 4)
    return summary


def markdown(results: dict, summary: dict) -> str:
    rows = []

    def row(measure: str, result) -> None:
        rows.append(f"| {measure} | {result} |")

    clean, labels = summary["clean_labels"], summary["labels"]
    row("Labels", f"{labels} ({clean} clean, {labels - clean} photo-like)")
    row("Errors", summary["errors"])
    row(
        "Clean labels with every field as expected",
        f"{summary['clean_labels_all_exact']} of {clean}",
    )
    if summary["photo_fields"]:
        row(
            "Photo-like label fields with an acceptable result",
            f"{summary['photo_fields_acceptable']} of {summary['photo_fields']}",
        )
    row("False passes (a planted problem shown as Matches)", len(summary["false_passes"]))
    row("False fails (a Problem where none was planted)", len(summary["false_fails"]))
    if "clean_fields_read" in summary:
        row(
            "Clean label fields read exactly as printed",
            f"{summary['clean_fields_read_exactly']} of {summary['clean_fields_read']}",
        )
    if "seconds" in summary:
        s = summary["seconds"]
        row("Time per label (median / p95 / max)", f"{s['median']} s / {s['p95']} s / {s['max']} s")
        row("Labels checked in under 5 seconds", f"{s['under_5_seconds']} of {s['timed']}")
    if "server_seconds" in summary:
        s = summary["server_seconds"]
        row(
            "Of that, time inside the app (median / p95 / max)",
            f"{s['median']} s / {s['p95']} s / {s['max']} s",
        )
    if "batch_seconds" in results:
        row("Whole batch, upload to finished", f"{results['batch_seconds']} s")
    if "outlined" in summary:
        row(
            "Fields to look at that were outlined on the picture",
            f"{summary['outlined']} of {summary['flagged_on_label']}",
        )
    if summary["read_by_ocr"]:
        row("Labels read with basic text reading", summary["read_by_ocr"])
    if "cost_usd" in summary:
        row("AI cost for this run (estimate from tokens)", f"${summary['cost_usd']}")

    if results["mode"] == "deployed app":
        mode = (
            f"Mode: the deployed app at {results['url']}, through its web pages. Times are "
            "single checks, one at a time, end to end from this computer: upload, reading, "
            "rules, outlines, and the page coming back. Field results are from one batch of "
            "all the labels."
        )
    else:
        mode = f"Mode: in process, reader `{results['reader']}`, settings `{results['settings']}`."
    kinds = ", ".join(
        f"{kind.replace('_', ' ')} {n}" for kind, n in sorted(summary["judgments"].items())
    )
    lines = [f"# Evaluation, {results['measured_at']}", "", mode, ""]
    if results["mode"] == "deployed app":
        version = f"Commit of this evaluation tool: `{results['commit']}`."
        if results.get("deployed_as"):
            version += f" The deployed app: {results['deployed_as']}."
    else:
        version = f"Commit: `{results['commit']}`."
    lines += [version, "", "## Summary", ""]
    lines += [
        "| Measure | Result |",
        "|---|---|",
        *rows,
        "",
        f"Field results by kind: {kinds}.",
        "",
    ]
    for title, items in (
        ("False passes", summary["false_passes"]),
        ("False fails", summary["false_fails"]),
    ):
        if items:
            lines += [f"## {title}", "", *[f"- {item}" for item in items], ""]
    lines += ["## By label", "", "| Label | Photo | Time | Result |", "|---|---|---|---|"]
    for record in results["records"]:
        outcome = record.get("error") or describe(record)
        seconds = f"{record['seconds']} s" if "seconds" in record else ""
        lines.append(f"| {record['id']} | {record['photo'] or ''} | {seconds} | {outcome} |")
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--url", help="evaluate a deployed app instead of running in this process")
    parser.add_argument("--reader", choices=["anthropic", "ocr"], default="anthropic")
    parser.add_argument("--only", help="comma-separated sample ids")
    parser.add_argument(
        "--deployed-as", help="what the deployed app runs, for the report: its revision and image"
    )
    args = parser.parse_args()

    # Taken before the run, so it is the code that ran. "-dirty" marks a run made with
    # changes that weren't committed yet.
    commit = subprocess.run(
        ["git", "describe", "--always", "--dirty"], capture_output=True, text=True, check=False
    ).stdout.strip()
    entries = load_entries(args.only.split(",") if args.only else None)
    if args.url:
        results = evaluate_deployed(entries, args.url)
        results["deployed_as"] = args.deployed_as
    else:
        results = asyncio.run(evaluate_here(entries, args.reader))
    results["measured_at"] = datetime.now().isoformat(timespec="seconds")
    results["commit"] = commit
    summary = summarize(results)
    results["summary"] = summary

    kind = "deployed" if args.url else args.reader
    stem = OUT_DIR / f"evaluation-{results['measured_at'][:10]}-{kind}"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stem.with_suffix(".json").write_text(json.dumps(results, indent=2, ensure_ascii=False) + "\n")
    stem.with_suffix(".md").write_text(markdown(results, summary))
    print()
    print(markdown(results, summary))
    print(f"Saved: {stem}.json and .md")


if __name__ == "__main__":
    main()
